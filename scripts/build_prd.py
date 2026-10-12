from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
OUTPUT_PATH = OUTPUT_DIR / "PRD_Sistem_KPI_PT_Berlian_Artha_Label.docx"

BLACK = "000000"
DARK = "27313A"
DARK_BLUE = "26384A"
PALE_BLUE = "EDF3F8"
PALE_GRAY = "F5F6F7"
MID_GRAY = "6B7280"
LIGHT_BORDER = "D9D9D9"
BAL_RED = "A61B1B"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_run_font(run, name="Aptos", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def configure_table(table, widths, header=True, compact=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        if header and row_idx == 0:
            set_repeat_table_header(row)
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                set_cell_width(cell, widths[idx])
            set_cell_border(cell)
            if compact:
                set_cell_margins(cell, top=70, bottom=70, start=85, end=85)
            else:
                set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths, *, compact=False, center_columns=None, keep_rows=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, header in enumerate(headers):
        set_cell_text(hdr.cells[i], header, bold=True, color="FFFFFF", size=9.3,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i], DARK_BLUE)
    for row_index, data in enumerate(rows):
        row = table.add_row()
        if keep_rows:
            keep_row_together(row)
        for col_index, value in enumerate(data):
            align = WD_ALIGN_PARAGRAPH.CENTER if center_columns and col_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[col_index], value, size=9.2 if compact else 9.5, align=align)
            if row_index % 2 == 1:
                set_cell_shading(row.cells[col_index], PALE_BLUE)
    configure_table(table, widths, header=True, compact=compact)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_kpi_catalog_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("Group", "Team", "Kandidat KPI")
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=9.1,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for row_index, (group, team, metrics) in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        set_cell_text(row.cells[0], group, bold=True, size=9.0)
        set_cell_text(row.cells[1], team, bold=True, size=9.0)
        cell = row.cells[2]
        cell.text = ""
        for metric_index, metric in enumerate(metrics):
            p = cell.paragraphs[0] if metric_index == 0 else cell.add_paragraph()
            p.style = doc.styles["List Bullet"]
            p.paragraph_format.left_indent = Inches(0.14)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.2)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(metric)
            set_run_font(r, size=8.9)
        if row_index % 2 == 1:
            for cell_item in row.cells:
                set_cell_shading(cell_item, PALE_BLUE)
    configure_table(table, [1.05, 1.30, 4.55], header=True, compact=True)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True)
    return p


def add_body(doc, text, *, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.widow_control = True
    if bold_lead:
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=10.8, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.8, italic=italic)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.28 + 0.18 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.widow_control = True
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.widow_control = True
        r = p.add_run(f"{index}.  {item}")
        set_run_font(r, size=10.5)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_specs = {
        "Heading 1": (16, 15, 7),
        "Heading 2": (13, 11, 5),
        "Heading 3": (11.5, 8, 4),
    }
    for style_name, (size, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for list_style in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[list_style]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.header_distance = Inches(0.34)
        section.footer_distance = Inches(0.34)
        section.different_first_page_header_footer = True


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run("PRD Sistem KPI PT Berlian Artha Label")
        set_run_font(hr, size=8.5, color=MID_GRAY)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.paragraph_format.space_before = Pt(0)
        left = fp.add_run("Version 0.1   Draft for Business Validation")
        set_run_font(left, size=8.5, color=MID_GRAY)
        tab_stops = fp.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.65))
        fp.add_run("\t")
        field_run = fp.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        value = OxmlElement("w:t")
        value.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field_run._r.extend([begin, instr, separate, value, end])
        set_run_font(field_run, size=8.5, color=BAL_RED)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("PT BERLIAN ARTHA LABEL")
    set_run_font(r, size=12, bold=True, color=BAL_RED)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(10)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    r = title.add_run("Product Requirements Document")
    set_run_font(r, name="Aptos Display", size=29, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(34)
    r = subtitle.add_run("Sistem Manajemen KPI PT Berlian Artha Label")
    set_run_font(r, name="Aptos Display", size=18, bold=True)

    meta_rows = [
        ("Versi", "0.1"),
        ("Status", "Draft untuk validasi bisnis"),
        ("Tanggal", "12 Oktober 2026"),
        ("Cakupan", "Company, group atau department, team, dan employee"),
        ("Unit operasional", "21 team"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for idx, (label, value) in enumerate(meta_rows):
        row = table.add_row()
        set_cell_text(row.cells[0], label, bold=True, size=10)
        set_cell_text(row.cells[1], value, size=10)
        set_cell_width(row.cells[0], 1.35)
        set_cell_width(row.cells[1], 4.8)
        for cell in row.cells:
            set_cell_margins(cell, top=85, bottom=85, start=0, end=90)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "5")
            bottom.set(qn("w:color"), LIGHT_BORDER)
            tc_borders.append(bottom)
            tc_pr.append(tc_borders)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Dokumen ini menetapkan kebutuhan produk untuk membangun sistem KPI yang terukur, dapat diaudit, dan mengikuti struktur organisasi PT Berlian Artha Label.")
    set_run_font(r, size=11.5)
    doc.add_page_break()


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_cover(doc)

    add_heading(doc, "Kontrol Dokumen", 1)
    add_table(
        doc,
        ("Versi", "Tanggal", "Status", "Keterangan"),
        [
            ("0.1", "12 Oktober 2026", "Draft untuk validasi bisnis", "Rancangan awal berdasarkan struktur organisasi dan kandidat KPI yang tersedia"),
        ],
        [0.7, 1.35, 1.75, 3.1],
        center_columns={0, 1, 2},
    )

    add_heading(doc, "Pihak yang Perlu Menyetujui", 2)
    add_table(
        doc,
        ("Pihak", "Fokus persetujuan", "Status"),
        [
            ("Management", "Tujuan, komposisi nilai, kewenangan approval, dan hasil akhir", "Belum disetujui"),
            ("HR Performance", "Periode, template KPI, kalibrasi, rating, dan tata kelola perubahan", "Belum disetujui"),
            ("Department Head atau Manager", "KPI, target, sumber data, dan jalur approval per unit", "Belum disetujui"),
            ("Team Leader", "Definisi operasional KPI, frekuensi input, bukti, dan pengecualian", "Belum disetujui"),
            ("IT atau Tim Pengembang", "Kelayakan teknis, integrasi, keamanan, dan dukungan operasional", "Belum ditinjau"),
        ],
        [1.5, 4.2, 1.2],
        center_columns={2},
    )

    add_heading(doc, "Ringkasan Produk", 1)
    add_body(
        doc,
        "Sistem KPI akan menggunakan struktur Company lalu Group atau Department lalu Team lalu Employee. "
        "Model ini mendukung 21 team operasional, termasuk tiga team Accounting dan sembilan team Produksi. "
        "Design A dan Design B diperlakukan sebagai dua team terpisah walaupun kandidat KPI awalnya sama."
    )
    add_body(
        doc,
        "Produk harus mengelola target, realisasi, bukti, verifikasi, review, kalibrasi, approval final, penguncian periode, dan koreksi yang tercatat. "
        "Setiap nilai harus dapat ditelusuri kembali ke definisi KPI, periode, sumber data, pihak yang memasukkan, pihak yang menyetujui, dan riwayat perubahannya."
    )
    add_body(
        doc,
        "Rilis pertama berfokus pada fondasi KPI yang berlaku lintas team. Fitur khusus seperti board PPIC, aging piutang, dan NCR atau CAPA dibangun di atas fondasi yang sama dan dapat dirilis bertahap."
    )

    add_heading(doc, "1 Tujuan dan Keputusan Produk", 1)
    add_heading(doc, "1 1 Tujuan", 2)
    add_bullets(doc, [
        "Menyediakan satu sumber data KPI perusahaan, group atau department, team, dan individu.",
        "Menstandarkan siklus penetapan target, pengisian realisasi, verifikasi, kalibrasi, approval, dan penguncian periode.",
        "Memberi Team Leader alat untuk memantau pencapaian team dan anggota serta menindaklanjuti masalah operasional.",
        "Menjaga integritas penilaian melalui kontrol akses, bukti pendukung, versioning, dan audit trail.",
        "Memberi Management dan HR dasar yang konsisten untuk review kinerja dan keputusan perbaikan.",
    ])

    add_heading(doc, "1 2 Keputusan Produk Utama", 2)
    add_table(
        doc,
        ("Area", "Keputusan"),
        [
            ("Struktur", "Company menjadi induk; group atau department menampung beberapa team; penilaian mencapai level team dan employee."),
            ("Kepemilikan", "Setiap team memiliki Team Leader. Jalur atasan di atas Team Leader dikonfigurasi per team."),
            ("Governance", "HR Performance mengelola periode, template, dan kalibrasi. Management memberi approval final."),
            ("Integritas data", "Data yang sudah disetujui tidak dapat ditimpa langsung. Koreksi harus melalui permintaan koreksi dengan alasan dan riwayat."),
            ("Konfigurasi", "Bobot, formula, rating, target, sumber data, dan jalur approval harus dapat dikonfigurasi dan diberi tanggal berlaku."),
            ("Pemisahan wewenang", "System Admin mengelola konfigurasi teknis dan akun, tetapi tidak mengubah nilai KPI."),
        ],
        [1.35, 5.55],
    )

    add_heading(doc, "1 3 Hasil yang Diharapkan", 2)
    add_bullets(doc, [
        "Seluruh KPI aktif mempunyai definisi, satuan, formula, arah pencapaian, target, frekuensi, bobot, sumber data, dan pemilik data.",
        "Management dapat melihat ringkasan perusahaan dan melakukan drill down sampai team serta individu sesuai kewenangannya.",
        "Team Leader dapat mengetahui KPI merah, anggota yang belum mengisi, data yang menunggu verifikasi, dan corrective action yang jatuh tempo.",
        "HR dapat menutup periode hanya setelah pemeriksaan kelengkapan, kalibrasi, dan approval selesai.",
        "Auditor dapat merekonstruksi perubahan nilai tanpa mendapatkan hak untuk mengeditnya.",
    ])

    add_heading(doc, "2 Latar Belakang dan Masalah", 1)
    add_body(
        doc,
        "Struktur organisasi PT Berlian Artha Label memiliki team dengan karakter KPI yang berbeda. Commercial dan Support lebih banyak menggunakan indikator layanan, penjualan, perencanaan, dan dukungan. Accounting memerlukan kontrol akurasi, ketepatan waktu, dan aging. Produksi memerlukan data output, yield, reject, waste, downtime, schedule attainment, serta kualitas per mesin atau shift."
    )
    add_body(
        doc,
        "Satu format penilaian yang kaku tidak cukup untuk seluruh unit. Sistem memerlukan template KPI per jabatan atau team, tetapi tetap menggunakan proses approval, aturan perhitungan, audit, dan struktur data yang sama. Tanpa definisi dan sumber data yang jelas, nilai dapat diperdebatkan dan sulit dibandingkan antarperiode."
    )
    add_body(
        doc,
        "PRD ini menetapkan kebutuhan produk. Penetapan target bisnis, angka bobot final, rating, formula rinci setiap KPI, dan integrasi ke sistem lain tetap memerlukan keputusan HR, Management, dan pemilik proses."
    )

    add_heading(doc, "3 Ruang Lingkup", 1)
    add_heading(doc, "3 1 Cakupan Rilis Inti", 2)
    add_bullets(doc, [
        "Master organisasi, karyawan, jabatan, team, atasan, shift, dan periode penilaian.",
        "Role based access untuk Management, Department Head atau Manager, Team Leader, Karyawan, HR Performance, Auditor, dan System Admin.",
        "Template KPI berversi per team atau jabatan, termasuk formula, arah pencapaian, bobot, target, sumber data, dan tanggal berlaku.",
        "Pengisian atau impor realisasi, lampiran bukti, validasi, komentar, dan status kelengkapan.",
        "Workflow approval yang dapat melewati Department Head jika posisi tersebut belum tersedia dan konfigurasi telah disetujui.",
        "Perhitungan skor individu dan team, kalibrasi, approval final, penguncian periode, serta permintaan koreksi.",
        "Dashboard sesuai peran, laporan scorecard, ekspor, notifikasi, dan audit trail.",
        "Halaman Team Leader dan daily team review."
    ])

    add_heading(doc, "3 2 Cakupan Bertahap Setelah Fondasi", 2)
    add_bullets(doc, [
        "Board status order untuk Admin PPIC.",
        "Aging dashboard untuk Account Receivable.",
        "Workflow NCR dan CAPA untuk Quality Control.",
        "Analisis produksi per shift, mesin, work order, defect, downtime, dan reason code.",
        "Integrasi otomatis dengan attendance, ERP atau production data, accounting, dan sumber lain setelah kontrak data disetujui.",
        "Analitik tren, perbandingan periode, dan rekomendasi improvement berbasis data tervalidasi."
    ])

    add_heading(doc, "3 3 Di Luar Cakupan Saat Ini", 2)
    add_bullets(doc, [
        "Perhitungan payroll, bonus, atau insentif otomatis berdasarkan skor KPI sebelum kebijakan kompensasi disetujui.",
        "Penggantian fungsi ERP, accounting, attendance, atau machine monitoring yang sudah ada.",
        "Perubahan struktur organisasi tanpa persetujuan HR dan Management.",
        "Penetapan target dan rating final oleh sistem tanpa keputusan pemilik proses.",
    ])

    add_heading(doc, "4 Struktur Organisasi di Sistem", 1)
    org_rows = [
        ("Commercial dan Support", "9", "Sales; Customer Service; R&D Costing; Design A; Design B; Admin PPIC; Purchasing; HR; GA"),
        ("Accounting", "3", "Admin Accounting; Account Receivable; Tax Invoice"),
        ("Produksi", "9", "Woven; Warehouse; Screen Print; Quality Control; Packing Logistic; Logistic; Hot Cut; Hang Tag; Flexo Printing"),
    ]
    add_table(doc, ("Group atau Department", "Jumlah team", "Team"), org_rows, [1.55, 0.9, 4.45], center_columns={1})
    add_body(doc, "Total unit operasional adalah 21 team. Accounting dan Produksi berfungsi sebagai kelompok induk. Setiap team mempunyai Team Leader dan dapat mempunyai Manager atau Department Head sebagai reviewer berikutnya.")

    add_heading(doc, "4 1 Aturan Struktur", 2)
    add_bullets(doc, [
        "Seorang karyawan memiliki satu home team aktif pada suatu tanggal, tetapi dapat menerima delegasi atau penugasan sementara yang tercatat.",
        "Perubahan team, Team Leader, Manager, shift, atau jabatan berlaku berdasarkan tanggal efektif dan tidak mengubah histori periode sebelumnya.",
        "Design A dan Design B disimpan sebagai team terpisah. Template KPI dapat sama, tetapi target, realisasi, scorecard, dan pencapaiannya tidak digabung.",
        "Jika Department Head belum ada, workflow dapat diarahkan dari Team Leader langsung ke HR Performance atau Management sesuai konfigurasi yang disetujui.",
        "Seorang approver tidak boleh menyetujui langkah yang sama atas data miliknya sendiri. Sistem menerapkan approver pengganti sesuai matriks kewenangan.",
    ])

    add_heading(doc, "5 Pengguna dan Hak Akses", 1)
    role_rows = [
        ("Management", "Melihat seluruh KPI perusahaan, memberi approval final, membuka ringkasan lintas unit", "Tidak mengubah realisasi sumber"),
        ("Department Head atau Manager", "Melihat dan mereview seluruh team di bawahnya", "Tidak melihat unit di luar kewenangan kecuali diberi akses"),
        ("Team Leader", "Mengelola target operasional yang didelegasikan, memverifikasi KPI team dan anggota, membuat daily review", "Tidak menimpa data yang sudah disetujui"),
        ("Karyawan", "Melihat target dan skor pribadi, mengisi realisasi yang menjadi tanggung jawabnya, mengunggah bukti", "Tidak melihat nilai individu lain"),
        ("HR Performance", "Mengelola periode, template, kalibrasi, rating, kelengkapan, dan workflow", "Perubahan nilai harus melalui proses dan tercatat"),
        ("Auditor", "Membaca hasil, definisi, bukti, approval, dan audit trail", "Read only"),
        ("System Admin", "Mengelola akun, role, parameter teknis, integrasi, dan health check", "Tidak mengubah target, realisasi, skor, atau approval"),
    ]
    add_table(doc, ("Peran", "Akses utama", "Batasan"), role_rows, [1.45, 3.85, 1.6], compact=True)

    add_heading(doc, "5 1 Prinsip Hak Akses", 2)
    add_bullets(doc, [
        "Hak akses mengikuti organisasi, peran, tanggal efektif, dan ruang lingkup data.",
        "Delegasi approval harus mempunyai pemberi delegasi, penerima, alasan, tanggal mulai, dan tanggal berakhir.",
        "Akses Auditor bersifat baca saja dan mencakup versi historis yang relevan.",
        "Akun nonaktif tidak dapat masuk, tetapi namanya tetap muncul pada histori transaksi.",
        "Aksi sensitif seperti membuka periode terkunci atau mengganti formula memerlukan kewenangan khusus dan alasan."
    ])

    add_heading(doc, "6 Alur Penilaian dan Status", 1)
    add_heading(doc, "6 1 Alur Penilaian Umum", 2)
    add_numbered(doc, [
        "HR Performance membuka periode dan menerbitkan template KPI yang berlaku.",
        "Target ditetapkan atau diimpor, lalu disetujui sesuai matriks kewenangan.",
        "Karyawan atau pemilik data memasukkan realisasi dan bukti sesuai frekuensi KPI.",
        "Team Leader memverifikasi kelengkapan, sumber, dan kewajaran data.",
        "Department Head atau Manager mereview hasil team jika peran tersebut tersedia.",
        "HR Performance melakukan pemeriksaan kelengkapan dan kalibrasi.",
        "Management memberi approval final.",
        "Sistem mengunci periode sehingga nilai tidak dapat diubah langsung."
    ])

    add_heading(doc, "6 2 Alur Validasi Produksi Harian", 2)
    add_numbered(doc, [
        "Operator atau Admin Produksi memasukkan atau mengimpor aktual per work order, shift, mesin, dan tanggal.",
        "Sistem memvalidasi format, duplikasi, referensi master, dan kelengkapan reason code.",
        "Team Leader memverifikasi aktual dan bukti sumber.",
        "Production Manager melakukan review bila dikonfigurasi.",
        "Data tervalidasi menjadi sumber perhitungan KPI harian dan agregasi periode."
    ])

    add_heading(doc, "6 3 Siklus Status", 2)
    status_rows = [
        ("Draft", "Data masih dapat diubah oleh pemilik"),
        ("Submitted", "Pemilik telah mengirim untuk verifikasi"),
        ("Verified", "Team Leader menerima data"),
        ("Reviewed", "Manager atau Department Head menyelesaikan review"),
        ("Calibrated", "HR menyelesaikan kalibrasi atau menyatakan tidak perlu kalibrasi"),
        ("Final Approved", "Management memberi approval final"),
        ("Locked", "Periode ditutup dan perubahan langsung ditolak"),
        ("Correction Requested", "Pengguna meminta koreksi dengan alasan dan bukti"),
        ("Reopened", "Data tertentu dibuka secara terbatas oleh pihak berwenang"),
    ]
    add_table(doc, ("Status", "Makna"), status_rows, [1.65, 5.25])

    add_heading(doc, "6 4 Koreksi dan Penolakan", 2)
    add_bullets(doc, [
        "Approver dapat mengembalikan data dengan komentar wajib. Status kembali ke pihak yang harus memperbaiki.",
        "Koreksi sesudah approval memerlukan permintaan koreksi yang menyebut nilai lama, nilai usulan, alasan, bukti, pemohon, dan approver.",
        "Sistem mempertahankan versi sebelum koreksi dan membuat versi baru setelah disetujui.",
        "Pembukaan periode tidak otomatis membuka semua data. Ruang lingkup koreksi dibatasi pada record yang disetujui untuk diperbaiki.",
    ])

    doc.add_page_break()
    add_heading(doc, "7 Model KPI dan Perhitungan Skor", 1)
    add_heading(doc, "7 1 Definisi KPI Wajib", 2)
    add_table(
        doc,
        ("Kelompok", "Field minimum"),
        [
            ("Identitas", "Kode KPI, nama, deskripsi, level penilaian, team atau jabatan, pemilik proses"),
            ("Pengukuran", "Satuan, frekuensi, arah KPI, formula, target, batas toleransi, pembulatan"),
            ("Penilaian", "Bobot, rating band, minimum score, maksimum score, aturan N/A"),
            ("Sumber", "Sistem atau dokumen sumber, pemilik data, metode input, bukti wajib"),
            ("Governance", "Versi, tanggal berlaku, approver, status aktif, alasan perubahan"),
        ],
        [1.25, 5.65],
    )

    add_heading(doc, "7 2 Arah dan Formula Dasar", 2)
    formula_rows = [
        ("Higher is better", "Nilai meningkat ketika aktual naik", "Aktual dibagi target dikali 100"),
        ("Lower is better", "Nilai meningkat ketika aktual turun", "Target dibagi aktual dikali 100 dengan aturan khusus untuk aktual nol"),
        ("Range", "Nilai terbaik berada pada rentang", "Skor ditentukan oleh rentang yang dikonfigurasi"),
        ("Binary", "Kondisi tercapai atau tidak", "Skor sesuai nilai Ya atau Tidak yang dikonfigurasi"),
        ("Milestone", "Pencapaian bertahap", "Skor mengikuti persentase atau tahap yang telah disetujui"),
    ]
    add_table(doc, ("Tipe", "Perilaku", "Rancangan formula"), formula_rows, [1.35, 2.35, 3.2])
    add_body(doc, "Formula pada tabel adalah rancangan awal. Aturan pembagian dengan nol, pembatasan score, toleransi, pembulatan, N/A, dan rating band harus ditetapkan sebelum konfigurasi digunakan untuk penilaian resmi.")

    add_heading(doc, "7 3 Komposisi Nilai Rancangan", 2)
    add_body(doc, "Komposisi berikut belum menjadi kebijakan final. Sistem harus memungkinkan HR dan Management mengubahnya melalui template berversi tanpa mengubah periode historis.")
    add_table(
        doc,
        ("Staff operasional", "Bobot awal"),
        [
            ("KPI individu", "50%"),
            ("KPI team", "30%"),
            ("Disiplin dan keselamatan", "10%"),
            ("Kompetensi dan perilaku", "10%"),
        ],
        [5.5, 1.4],
        center_columns={1},
    )
    add_table(
        doc,
        ("Team Leader", "Bobot awal"),
        [
            ("Pencapaian KPI team", "55%"),
            ("Kualitas dan produktivitas", "15%"),
            ("Leadership dan pengelolaan anggota", "15%"),
            ("Improvement", "10%"),
            ("Keselamatan dan kedisiplinan", "5%"),
        ],
        [5.5, 1.4],
        center_columns={1},
    )

    add_heading(doc, "7 4 Aturan Perhitungan", 2)
    add_bullets(doc, [
        "Total bobot dalam satu scorecard harus 100 persen sebelum template dapat diterbitkan.",
        "Skor team dan individu dihitung dari data yang berstatus memenuhi syarat sesuai konfigurasi periode.",
        "KPI tanpa data tidak otomatis diberi nol atau dikeluarkan. Sistem mengikuti aturan missing data yang disetujui untuk KPI tersebut.",
        "Perubahan formula, bobot, atau rating membuat versi baru dan hanya berlaku untuk periode yang belum dikunci atau periode mendatang.",
        "Sistem menampilkan rincian perhitungan sehingga pengguna berwenang dapat merekonsiliasi target, aktual, skor KPI, bobot, dan kontribusi ke total.",
        "Skor hasil kalibrasi harus menyimpan nilai sebelum kalibrasi, nilai setelah kalibrasi, alasan, dan approver."
    ])

    add_heading(doc, "8 Kebutuhan Fungsional", 1)
    fr_rows = [
        ("FR 001", "Organisasi dan histori", "Must", "Sistem menyimpan company, group atau department, team, employee, atasan, shift, jabatan, dan tanggal efektif tanpa mengubah histori."),
        ("FR 002", "Akses berbasis peran dan cakupan", "Must", "Setiap peran hanya dapat melihat dan melakukan aksi pada data yang diizinkan oleh organisasi dan workflow."),
        ("FR 003", "Periode penilaian", "Must", "HR dapat membuat, membuka, menutup, dan mengunci periode dengan validasi kelengkapan dan approval."),
        ("FR 004", "Template KPI berversi", "Must", "Template menyimpan definisi, formula, bobot, sumber, frekuensi, approver, dan tanggal berlaku."),
        ("FR 005", "Penetapan target", "Must", "Target dapat diinput atau diimpor, divalidasi, disetujui, dan ditelusuri per employee atau team."),
        ("FR 006", "Realisasi dan bukti", "Must", "Pengguna dapat mengisi atau mengimpor aktual serta bukti; sistem menandai data tidak lengkap atau tidak valid."),
        ("FR 007", "Approval configurable", "Must", "Workflow mengikuti Team Leader, Manager atau Head jika ada, HR, dan Management sesuai konfigurasi."),
        ("FR 008", "Koreksi terkendali", "Must", "Data disetujui tidak dapat ditimpa; koreksi menghasilkan request, approval, dan versi baru."),
        ("FR 009", "Perhitungan transparan", "Must", "Sistem menghitung skor sesuai versi formula dan menampilkan breakdown yang dapat direkonsiliasi."),
        ("FR 010", "Kalibrasi", "Must", "HR dapat menjalankan kalibrasi dengan alasan dan histori nilai sebelum serta sesudah."),
        ("FR 011", "Dashboard per peran", "Must", "Dashboard menampilkan ringkasan, exception, tren, dan drill down sesuai kewenangan."),
        ("FR 012", "Halaman Team Leader", "Must", "Team Leader melihat anggota, target, aktual, pending verification, KPI merah, masalah, tindakan, dan ekspor."),
        ("FR 013", "Daily team review", "Must", "Team Leader dapat mencatat masalah, penyebab, dampak, corrective action, PIC, deadline, status, dan bukti."),
        ("FR 014", "Notifikasi dan reminder", "Should", "Sistem mengirim notifikasi atas tugas, tenggat, penolakan, koreksi, dan keterlambatan."),
        ("FR 015", "Laporan dan ekspor", "Must", "Pengguna dapat mengekspor scorecard dan daftar exception sesuai hak akses dengan metadata periode dan versi."),
        ("FR 016", "Audit trail", "Must", "Aksi create, update, submit, approve, reject, calibrate, reopen, export, dan login sensitif tercatat."),
        ("FR 017", "Delegasi", "Should", "Approver dapat mendelegasikan dalam periode waktu terbatas dengan alasan dan histori."),
        ("FR 018", "Impor massal", "Should", "Impor menyediakan validasi per baris, preview, hasil diterima atau ditolak, dan pencegahan duplikasi."),
        ("FR 019", "Fitur team khusus", "Should", "PPIC board, AR aging, dan QC NCR atau CAPA menggunakan master serta kontrol akses yang sama."),
        ("FR 020", "System administration", "Must", "Admin mengelola akun dan konfigurasi teknis tanpa memperoleh izin mengubah nilai KPI."),
    ]
    add_table(doc, ("ID", "Kebutuhan", "Prioritas", "Kriteria penerimaan ringkas"), fr_rows, [0.65, 1.55, 0.75, 3.95], compact=True, center_columns={0, 2}, keep_rows=True)

    add_heading(doc, "8 1 Validasi Data", 2)
    add_bullets(doc, [
        "Sistem menolak referensi employee, team, mesin, work order, produk, material, customer, defect, downtime, atau reason code yang tidak aktif pada tanggal transaksi.",
        "Sistem mendeteksi duplikasi berdasarkan key sumber yang disetujui, misalnya tanggal, work order, employee, shift, mesin, dan jenis KPI.",
        "Nilai di luar batas wajar diberi flag dan memerlukan komentar atau approval tambahan sesuai konfigurasi.",
        "File impor menghasilkan ringkasan jumlah baris diterima, ditolak, diperbarui, dan duplikat serta menyediakan detail alasan per baris.",
        "Data manual menyimpan pembuat, waktu, sumber, dan bukti bila KPI mewajibkannya."
    ])

    add_heading(doc, "9 Dashboard dan Pengalaman Pengguna", 1)
    dashboard_rows = [
        ("Management", "Skor perusahaan dan group, tren, team merah, approval final, isu lintas fungsi, drill down"),
        ("Manager atau Department Head", "Ringkasan team, perbandingan team, pending review, exception, action overdue"),
        ("Team Leader", "Skor team, anggota, target versus aktual, pending verification, missing actual, KPI merah, shift, action"),
        ("Karyawan", "Target pribadi, aktual, status submission, feedback, skor, histori"),
        ("HR Performance", "Kesiapan periode, kelengkapan template, progress review, kalibrasi, distribusi rating, exception"),
        ("Auditor", "Filter periode dan unit, hasil final, bukti, approval chain, perubahan nilai"),
        ("System Admin", "Status akun, integrasi, job impor, error log, konfigurasi teknis"),
    ]
    add_table(doc, ("Pengguna", "Informasi utama"), dashboard_rows, [1.55, 5.35])

    add_heading(doc, "9 1 Halaman Team Leader", 2)
    add_bullets(doc, [
        "Ringkasan KPI team dan daftar anggota.",
        "Target serta aktual harian, mingguan, dan bulanan.",
        "Pending verification dan anggota yang belum mengisi realisasi.",
        "KPI berstatus merah serta perbandingan target dan aktual.",
        "Perbandingan antarshift jika team menggunakan shift.",
        "Daftar masalah, penyebab, corrective action, PIC, deadline, status, dan bukti.",
        "Coaching notes dengan kontrol akses yang sesuai.",
        "Permintaan revisi data, approval cuti atau delegasi penilaian, dan riwayat perubahan anggota.",
        "Ekspor scorecard team dengan periode, versi template, dan waktu ekspor."
    ])

    add_heading(doc, "9 2 Daily Team Review", 2)
    add_table(
        doc,
        ("Field", "Aturan"),
        [
            ("Masalah hari ini", "Wajib dan dapat dikaitkan ke KPI, work order, mesin, shift, atau customer"),
            ("Penyebab", "Dapat dipilih dari reason code dan diberi catatan"),
            ("Dampak terhadap KPI", "Memilih satu atau lebih KPI yang terdampak"),
            ("Corrective action", "Menjelaskan tindakan yang akan dilakukan"),
            ("PIC dan deadline", "Wajib sebelum action disimpan sebagai Open"),
            ("Status", "Open, In Progress, Waiting, Verified, atau Closed"),
            ("Bukti penyelesaian", "Wajib sebelum verifikasi Closed bila dikonfigurasi"),
        ],
        [1.7, 5.2],
    )

    add_heading(doc, "10 Fitur Khusus Unit", 1)
    add_heading(doc, "10 1 Board Admin PPIC", 2)
    add_body(doc, "Board menampilkan order berdasarkan status Waiting Material, Waiting Design, Ready for Production, In Production, QC, Packing, Ready to Deliver, dan Completed. Setiap perpindahan status menyimpan waktu, pengguna, alasan bila terlambat, serta referensi work order.")
    add_bullets(doc, [
        "Filter berdasarkan customer, order, work order, tanggal kebutuhan, mesin, dan status.",
        "Penanda work order overdue, material belum siap, perubahan jadwal, dan emergency job.",
        "Perbandingan plan versus actual dan waktu tunggu antarstatus.",
        "Audit trail untuk perubahan jadwal dan handover antarproses."
    ])

    add_heading(doc, "10 2 Aging Account Receivable", 2)
    add_body(doc, "Dashboard mengelompokkan invoice menjadi belum jatuh tempo, 1 sampai 30 hari, 31 sampai 60 hari, 61 sampai 90 hari, dan lebih dari 90 hari. Pengguna dapat menelusuri customer, invoice, nilai outstanding, dispute, tindak lanjut terakhir, dan perkiraan collection.")

    add_heading(doc, "10 3 NCR dan CAPA Quality Control", 2)
    add_body(doc, "Workflow QC mencakup Temuan, Containment, Root Cause, Corrective Action, Verification, dan Closed. Setiap tahap mempunyai PIC, deadline, bukti, dan approval sesuai konfigurasi. Repeat defect harus dapat dikaitkan ke NCR atau CAPA sebelumnya.")

    add_heading(doc, "11 Data dan Model Informasi", 1)
    entity_rows = [
        ("OrganizationUnit", "Company, group atau department, dan team dengan hubungan induk serta tanggal efektif"),
        ("EmployeeAssignment", "Employee, NIK, jabatan, team, leader, manager, shift, dan periode berlaku"),
        ("KPIDefinition", "Definisi KPI, formula, satuan, arah, frekuensi, sumber, dan versi"),
        ("KPITemplate", "Kumpulan KPI dan bobot untuk team atau jabatan pada periode tertentu"),
        ("Target", "Nilai target per KPI, employee atau team, periode, dan approval"),
        ("Actual", "Nilai aktual, dimensi operasional, sumber, bukti, status, dan versi"),
        ("Score", "Hasil formula, bobot, kontribusi, rating, dan versi perhitungan"),
        ("ApprovalTask", "Tahap, approver, keputusan, komentar, waktu, dan delegasi"),
        ("CorrectionRequest", "Nilai lama, usulan, alasan, bukti, keputusan, dan versi baru"),
        ("TeamReview", "Masalah, penyebab, dampak, corrective action, PIC, deadline, status, dan bukti"),
        ("AuditEvent", "Aktor, aksi, waktu, objek, nilai sebelum dan sesudah, serta metadata perangkat atau sumber"),
    ]
    add_table(doc, ("Entitas", "Isi utama"), entity_rows, [1.75, 5.15])

    add_heading(doc, "11 1 Data Master", 2)
    add_bullets(doc, [
        "Karyawan, nomor induk karyawan, jabatan, team, Team Leader, Manager atau Head, dan shift.",
        "Mesin, proses produksi, produk atau SKU, customer, material, dan work order.",
        "Defect code, downtime code, reason code, lokasi warehouse, dan status operasional.",
        "Periode penilaian, template KPI per jabatan atau team, serta sumber data setiap KPI.",
        "Kalender kerja dan tanggal efektif untuk target, assignment, serta formula."
    ])

    add_heading(doc, "11 2 Sumber Data", 2)
    source_rows = [
        ("Input manual", "Digunakan bila belum ada sumber sistem; bukti dan audit wajib sesuai konfigurasi"),
        ("Impor file", "Menggunakan template, validasi, preview, rejection report, dan idempotency key"),
        ("Integrasi API atau database", "Menggunakan mapping, monitoring, retry, serta rekonsiliasi sumber dan target"),
        ("Perhitungan sistem", "Dihasilkan dari data tervalidasi menggunakan formula berversi"),
    ]
    add_table(doc, ("Metode", "Kontrol minimum"), source_rows, [1.6, 5.3])

    add_heading(doc, "12 Notifikasi dan Eskalasi", 1)
    add_bullets(doc, [
        "Pengingat sebelum tenggat pengisian, verifikasi, review, kalibrasi, dan approval.",
        "Notifikasi saat data ditolak, dikembalikan, diminta koreksi, atau dibuka kembali.",
        "Eskalasi tugas yang melewati SLA kepada atasan berikutnya sesuai konfigurasi.",
        "Notifikasi corrective action yang mendekati deadline atau overdue.",
        "Ringkasan periodik untuk Team Leader, Manager, HR, dan Management sesuai preferensi dan kewenangan.",
        "Tautan notifikasi membuka record yang tepat setelah pengguna melewati pemeriksaan akses."
    ])

    add_heading(doc, "13 Laporan dan Ekspor", 1)
    add_bullets(doc, [
        "Scorecard individu, team, group atau department, dan perusahaan.",
        "Target versus aktual, tren periode, distribusi rating, dan daftar KPI merah.",
        "Kelengkapan input, pending approval, missing data, dan overdue action.",
        "Perbandingan team atau shift hanya untuk pengguna yang mempunyai ruang lingkup akses.",
        "Audit report berisi definisi KPI, versi, perubahan nilai, keputusan approval, dan koreksi.",
        "Ekspor mencantumkan nama laporan, periode, filter, waktu pembuatan, pembuat, dan versi template."
    ])

    add_heading(doc, "14 Audit Keamanan dan Kepatuhan", 1)
    audit_rows = [
        ("Access control", "Least privilege, pemisahan peran, cakupan organisasi, dan pemeriksaan akses pada setiap aksi"),
        ("Authentication", "Kebijakan login dan session mengikuti standar IT perusahaan; dukungan SSO dapat ditambahkan bila tersedia"),
        ("Audit trail", "Append only secara logis; mencatat nilai sebelum dan sesudah serta tidak dapat diubah oleh pengguna biasa"),
        ("Data protection", "Transmisi terenkripsi, perlindungan data personal, serta pembatasan akses ke coaching notes dan bukti sensitif"),
        ("Period lock", "Record final tidak dapat diubah tanpa proses reopen atau correction yang sah"),
        ("Segregation of duties", "System Admin tidak dapat mengubah nilai; approver tidak menyetujui langkah yang sama untuk datanya sendiri"),
        ("Export control", "Ekspor mengikuti ruang lingkup akses dan dicatat bila berisi data sensitif"),
        ("Retention", "Masa simpan data, bukti, dan audit ditetapkan HR, Legal, dan IT sebelum produksi"),
    ]
    add_table(doc, ("Area", "Kebutuhan"), audit_rows, [1.55, 5.35])

    add_heading(doc, "15 Kebutuhan Nonfungsional", 1)
    nfr_rows = [
        ("NFR 001", "Keandalan transaksi", "Submit, approval, koreksi, dan perhitungan tidak boleh menghasilkan status parsial."),
        ("NFR 002", "Kinerja", "Halaman standar ditargetkan tampil dalam 3 detik pada beban yang disepakati; target diuji ulang setelah jumlah pengguna dan volume data tersedia."),
        ("NFR 003", "Ketersediaan", "Sistem tersedia pada jam operasional; SLA numerik dan maintenance window ditetapkan bersama IT."),
        ("NFR 004", "Waktu", "Timestamp disimpan konsisten dan ditampilkan dalam zona waktu Asia Jakarta."),
        ("NFR 005", "Usability", "Tabel dapat difilter, diurutkan, dicari, dan diekspor; status dan tindakan utama terlihat jelas."),
        ("NFR 006", "Responsive", "Fungsi utama dapat digunakan pada desktop dan tablet; kebutuhan mobile untuk operator divalidasi dalam pilot."),
        ("NFR 007", "Aksesibilitas", "Navigasi keyboard, label form, kontras, fokus, dan pesan error dapat dipahami tanpa bergantung pada warna saja."),
        ("NFR 008", "Observability", "Job impor, integrasi, error perhitungan, dan aktivitas sensitif dapat dipantau dan diberi alert."),
        ("NFR 009", "Backup dan recovery", "RPO, RTO, frekuensi backup, dan uji restore ditetapkan IT sebelum go live."),
        ("NFR 010", "Kompatibilitas", "Dukungan browser perusahaan ditetapkan dan diuji sebelum UAT."),
    ]
    add_table(doc, ("ID", "Area", "Kebutuhan"), nfr_rows, [0.75, 1.4, 4.75], compact=True, center_columns={0})

    add_heading(doc, "16 Kriteria Penerimaan Rilis Inti", 1)
    acceptance = [
        "Struktur tiga group atau department dan 21 team tersedia dengan hubungan atasan serta tanggal efektif yang benar.",
        "Pengujian hak akses membuktikan bahwa setiap peran hanya melihat dan menjalankan aksi yang diizinkan.",
        "HR dapat membuat periode dan menerbitkan template KPI berversi dengan total bobot 100 persen.",
        "Target dan realisasi dapat dimasukkan atau diimpor, divalidasi, dilengkapi bukti, dan diproses sampai approval final.",
        "Perhitungan scorecard dapat direkonsiliasi dari target, aktual, formula, skor KPI, bobot, dan kontribusi total.",
        "Data final tidak dapat ditimpa langsung dan koreksi yang disetujui menghasilkan versi baru serta audit lengkap.",
        "Dashboard Team Leader menampilkan anggota, pending verification, missing actual, KPI merah, masalah, dan corrective action.",
        "HR dapat melakukan kalibrasi dan periode hanya dapat dikunci setelah pemeriksaan kelengkapan dan approval.",
        "Management dapat melihat ringkasan perusahaan serta drill down sesuai hak akses.",
        "Auditor dapat menelusuri definisi, sumber, bukti, approval chain, dan perubahan tanpa hak edit.",
        "Ekspor mengikuti filter dan hak akses serta mencantumkan metadata periode dan versi.",
        "UAT bisnis disetujui oleh Management, HR Performance, dan perwakilan unit pilot tanpa defect kritis yang terbuka."
    ]
    add_numbered(doc, acceptance)

    add_heading(doc, "17 Rencana Implementasi Bertahap", 1)
    phase_rows = [
        ("Tahap 0", "Validasi governance dan data", "Struktur, approver, bobot, rating, formula, KPI owner, sumber data, master, dan kebijakan koreksi disetujui"),
        ("Tahap 1", "Fondasi produk", "Organisasi, akun, periode, template, target, aktual, bukti, approval, scoring, audit, dashboard dasar, dan ekspor"),
        ("Tahap 2", "Pilot lintas fungsi", "Minimal satu team Commercial dan Support, satu Accounting, dan satu Produksi; team final dipilih Management"),
        ("Tahap 3", "Rollout 21 team", "Template dan data seluruh team, pelatihan, migrasi, UAT per group, serta support go live"),
        ("Tahap 4", "Fitur khusus dan integrasi", "PPIC board, AR aging, QC NCR atau CAPA, produksi per shift atau mesin, dan integrasi sumber"),
        ("Tahap 5", "Optimasi", "Analitik tren, kualitas data, penyempurnaan SLA, dan review efektivitas KPI"),
    ]
    add_table(doc, ("Tahap", "Fokus", "Exit criteria"), phase_rows, [0.75, 1.55, 4.6], compact=True, center_columns={0})

    add_heading(doc, "18 Risiko dan Mitigasi", 1)
    risk_rows = [
        ("Definisi KPI berbeda antarunit", "Tinggi", "Workshop KPI dictionary dan sign off pemilik proses sebelum konfigurasi"),
        ("Sumber data tidak konsisten", "Tinggi", "Data profiling, pemilik data, validation rule, rekonsiliasi, dan pilot"),
        ("Jalur approval belum pasti", "Tinggi", "Tetapkan approver per team dan rule pengganti sebelum UAT"),
        ("Bobot atau rating belum disetujui", "Tinggi", "Simpan sebagai konfigurasi draft; blok penggunaan resmi sebelum approval"),
        ("Perubahan organisasi merusak histori", "Sedang", "Gunakan assignment dan hierarchy bertanggal efektif"),
        ("Input manual terlambat", "Sedang", "Reminder, escalation, completeness dashboard, dan ownership yang jelas"),
        ("KPI mendorong perilaku yang salah", "Sedang", "Review lintas fungsi, balance quality dan quantity, serta evaluasi pascapilot"),
        ("Akses berlebihan", "Tinggi", "Role matrix, least privilege, access test, dan review berkala"),
        ("Integrasi menghasilkan duplikasi", "Sedang", "Idempotency key, source reference, reconciliation, dan retry terkontrol"),
        ("Periode dibuka tanpa kontrol", "Tinggi", "Reopen terbatas, alasan wajib, approval, dan audit"),
    ]
    add_table(doc, ("Risiko", "Dampak", "Mitigasi"), risk_rows, [2.25, 0.8, 3.85], compact=True, center_columns={1})

    add_heading(doc, "19 Keputusan yang Masih Terbuka", 1)
    decisions = [
        "Apa perbedaan tanggung jawab, proses, produk, atau customer antara Design A dan Design B?",
        "Siapa atasan dan approver di atas setiap Team Leader, termasuk jalur pengganti saat posisi kosong?",
        "Apakah bobot staff operasional dan Team Leader disetujui atau perlu dibedakan per jabatan dan team?",
        "Apa rating band, score cap, aturan N/A, aturan missing data, toleransi, dan pembulatan yang berlaku?",
        "KPI mana yang wajib pada rilis inti dan KPI mana yang hanya kandidat untuk tahap berikutnya?",
        "Apa target, formula, frekuensi, bukti, pemilik data, dan sistem sumber untuk setiap KPI?",
        "Team mana yang menjadi pilot dan berapa jumlah pengguna serta volume data hariannya?",
        "Sistem apa yang akan diintegrasikan, metode aksesnya, dan siapa pemilik teknis serta bisnisnya?",
        "Berapa SLA approval, correction, corrective action, dan notification escalation?",
        "Berapa masa simpan data, RPO, RTO, SLA availability, browser yang didukung, dan kebijakan akses data personal?",
        "Apakah coaching notes termasuk data HR terbatas dan siapa yang dapat membacanya?",
        "Apakah skor KPI akan digunakan untuk bonus atau keputusan kompensasi? Jika ya, diperlukan kebijakan terpisah yang disetujui."
    ]
    add_numbered(doc, decisions)

    doc.add_page_break()
    add_heading(doc, "Lampiran A Kandidat KPI per Team", 1)
    add_body(doc, "Daftar berikut merupakan kandidat awal. Setiap KPI baru dapat diaktifkan setelah definisi, formula, target, sumber data, pemilik, frekuensi, bobot, dan approver disetujui. Design A dan Design B ditampilkan sebagai team terpisah walaupun kandidat KPI awalnya sama.")

    kpi_catalog = [
        ("Commercial dan Support", "Sales", ["Revenue achievement", "Sales volume", "Gross margin", "Customer baru", "Repeat order", "Customer retention", "Quotation conversion rate", "Forecast accuracy", "Jumlah order batal", "Kelengkapan informasi order", "Jumlah complaint per customer", "Collection support untuk invoice bermasalah"]),
        ("Commercial dan Support", "Customer Service", ["Response time kepada customer", "Order entry accuracy", "Kelengkapan purchase order dan spesifikasi", "Kecepatan konfirmasi order", "Order handover ke PPIC tepat waktu", "Jumlah revisi akibat kesalahan input", "Complaint response time", "Complaint closure time", "Customer satisfaction", "Open request yang melewati SLA"]),
        ("Commercial dan Support", "R&D Costing", ["Costing lead time", "Ketepatan estimasi biaya", "Selisih estimated cost versus actual cost", "Trial success rate", "Sample completion lead time", "Material usage optimization", "Jumlah revisi costing", "Kecepatan pembuatan spesifikasi", "Produk baru selesai tepat waktu", "Cost reduction dari improvement", "Kelengkapan BOM dan routing", "Handover ke produksi tanpa masalah"]),
        ("Commercial dan Support", "Design A", ["Artwork lead time", "First approval pass rate", "Jumlah revisi internal", "Jumlah revisi dari customer", "Kesalahan artwork yang masuk produksi", "Ketepatan teks ukuran warna dan barcode", "File selesai sesuai jadwal", "Kepatuhan terhadap versi artwork", "Design backlog", "Rework akibat kesalahan design", "Produktivitas design per orang", "Approval customer turnaround"]),
        ("Commercial dan Support", "Design B", ["Artwork lead time", "First approval pass rate", "Jumlah revisi internal", "Jumlah revisi dari customer", "Kesalahan artwork yang masuk produksi", "Ketepatan teks ukuran warna dan barcode", "File selesai sesuai jadwal", "Kepatuhan terhadap versi artwork", "Design backlog", "Rework akibat kesalahan design", "Produktivitas design per orang", "Approval customer turnaround"]),
        ("Commercial dan Support", "Admin PPIC", ["Production schedule attainment", "Work order release tepat waktu", "Plan versus actual production", "Material readiness", "Jumlah perubahan jadwal", "Backlog order", "Work order overdue", "Keakuratan jadwal mesin", "Kelengkapan dokumen produksi", "Lead time perencanaan", "Emergency job frequency", "Koordinasi handover antarproses", "Ketepatan update status order"]),
        ("Commercial dan Support", "Purchasing", ["Supplier on time delivery", "Incoming acceptance rate", "Purchase request to purchase order lead time", "Material shortage", "Emergency purchase", "Purchase price variance", "Purchase order accuracy", "Supplier complaint closure", "Supplier response time", "Approved supplier compliance", "Material datang sesuai kebutuhan PPIC", "Supplier performance score"]),
        ("Commercial dan Support", "HR", ["Recruitment lead time", "Kehadiran karyawan", "Ketepatan data absensi", "Employee turnover", "Training completion", "Competency fulfillment", "Performance review completion", "Masa probation selesai tepat waktu", "Kelengkapan dokumen karyawan", "Employee complaint resolution", "Disciplinary case closure", "Manpower plan versus actual"]),
        ("Commercial dan Support", "GA", ["Penyelesaian permintaan fasilitas sesuai SLA", "Facility uptime", "Kepatuhan 5R atau 5S", "Utility consumption", "Pengelolaan kendaraan", "Kepatuhan perizinan", "Pengelolaan keamanan", "Kebersihan area kerja", "General service request completion", "Pengeluaran GA versus anggaran", "Safety finding closure", "Ketersediaan fasilitas produksi"]),
        ("Accounting", "Admin Accounting", ["Ketepatan input jurnal", "Closing bulanan tepat waktu", "Reconciliation completion", "Document completeness", "Posting accuracy", "Jumlah jurnal koreksi", "Bank reconciliation accuracy", "Expense verification lead time", "Financial report timeliness", "Open reconciliation items", "Audit finding closure", "Kelengkapan dokumen transaksi"]),
        ("Accounting", "Account Receivable", ["Days Sales Outstanding", "Collection rate", "Overdue invoice percentage", "Aging receivable", "Invoice dispute resolution time", "Customer payment allocation accuracy", "Unapplied payment", "Collection follow up completion", "Bad debt exposure", "Customer statement accuracy", "Outstanding invoice tanpa tindak lanjut", "Collection forecast accuracy"]),
        ("Accounting", "Tax Invoice", ["Faktur pajak diterbitkan tepat waktu", "Tax invoice accuracy", "Jumlah faktur pajak yang direvisi", "Reconciliation invoice versus faktur pajak", "Kelengkapan data NPWP atau NIK customer", "Kepatuhan deadline pelaporan", "Rejected tax invoice", "Penyelesaian dokumen pajak", "Outstanding faktur pajak", "Audit atau tax finding closure"]),
        ("Produksi", "Woven", ["Good output", "Output per jam", "Production yield", "Waste percentage", "Reject rate", "Rework rate", "Machine utilization", "Downtime", "Schedule attainment", "Setup atau changeover time", "Productivity per operator", "Kesalahan spesifikasi produk"]),
        ("Produksi", "Warehouse", ["Inventory accuracy", "Picking accuracy", "Receiving accuracy", "Material issue lead time", "Stockout frequency", "FIFO atau FEFO compliance", "Selisih stock opname", "Damaged material", "Material traceability", "Slow moving stock", "Finished goods accuracy", "Kesalahan lot material"]),
        ("Produksi", "Screen Print", ["Good output per jam", "First pass yield", "Reject printing", "Rework", "Setup time", "Changeover time", "Waste material", "Machine availability", "Schedule attainment", "Downtime", "Defect berdasarkan jenis", "Kepatuhan parameter produksi"]),
        ("Produksi", "Quality Control", ["Incoming rejection", "In process rejection", "Final rejection", "First pass yield", "Inspection lead time", "Customer complaint rate", "Customer return rate", "Repeat defect", "Corrective action closure", "Release produk tepat waktu", "Calibration compliance", "Audit finding closure"]),
        ("Produksi", "Packing Logistic", ["Packing accuracy", "Packing output per jam", "Salah jumlah packing", "Salah label atau karton", "Damage saat packing", "Packing lead time", "Finished goods handover tepat waktu", "Packing material waste", "Repacking rate", "Kelengkapan dokumen", "Schedule attainment", "Produk tertahan di area packing"]),
        ("Produksi", "Logistic", ["On Time Delivery", "On Time In Full", "Delivery accuracy", "Proof of delivery completion", "Damage during delivery", "Salah customer atau alamat", "Return akibat kesalahan delivery", "Delivery cost", "Transporter performance", "Vehicle utilization", "Delivery document completeness", "Delivery complaint closure"]),
        ("Produksi", "Hot Cut", ["Output per jam", "Good quantity", "Reject cutting", "Kesalahan ukuran", "Waste", "Setup time", "Changeover time", "Machine downtime", "Schedule attainment", "Rework", "Productivity per operator", "Kepatuhan standar proses"]),
        ("Produksi", "Hang Tag", ["Output per jam", "First pass yield", "Reject rate", "Rework rate", "Kesalahan ukuran dan finishing", "Schedule attainment", "Setup time", "Waste material", "Machine utilization", "Order completion accuracy", "Packing accuracy", "Productivity per operator"]),
        ("Produksi", "Flexo Printing", ["OEE", "Good output", "Production speed versus standard", "Substrate waste", "Ink waste", "Setup waste", "Register defect", "Color defect", "Printing defect", "Changeover time", "Downtime", "Schedule attainment", "First pass yield", "Rework", "Produktivitas per shift"]),
    ]
    add_kpi_catalog_table(doc, kpi_catalog)

    add_heading(doc, "Lampiran B Aturan Khusus Kandidat KPI", 1)
    add_bullets(doc, [
        "Costing harus dapat dibandingkan dengan actual production untuk mengukur akurasi estimasi material, waktu mesin, dan tenaga kerja.",
        "Revisi Design harus mempunyai reason category: kesalahan internal, perubahan customer, perubahan Sales, atau masalah produksi. KPI Design tidak boleh membebankan revisi eksternal sebagai kesalahan internal.",
        "Warehouse membedakan lokasi Raw Material, Work in Process, Finished Goods, Reject atau Quarantine, Return, dan Scrap.",
        "Packing Logistic dan Logistic tetap menjadi dua team. Packing Logistic menyiapkan barang sebelum dikirim, sedangkan Logistic mengirim dan memastikan barang diterima.",
        "Data produksi harus mendukung dimensi tanggal, shift, mesin, operator, work order, produk, defect, downtime, dan reason code bila relevan.",
        "Dashboard Sales memerlukan target dan realisasi per sales, pipeline quotation, customer pasif, order margin rendah, forecast versus actual, dan complaint per sales atau customer."
    ])

    add_heading(doc, "Lampiran C Glosarium", 1)
    glossary_rows = [
        ("Actual", "Nilai realisasi KPI pada periode atau frekuensi tertentu"),
        ("CAPA", "Corrective and Preventive Action"),
        ("Correction Request", "Permintaan resmi untuk memperbaiki data yang telah melewati approval"),
        ("Drill down", "Navigasi dari ringkasan menuju rincian group, team, employee, atau transaksi sumber"),
        ("KPI", "Key Performance Indicator"),
        ("NCR", "Non Conformance Report"),
        ("OEE", "Overall Equipment Effectiveness"),
        ("OTIF", "On Time In Full"),
        ("Period Lock", "Kondisi ketika data final tidak dapat diubah langsung"),
        ("Scorecard", "Kumpulan KPI, target, aktual, skor, bobot, dan hasil akhir untuk objek penilaian"),
        ("Template KPI", "Definisi KPI dan aturan penilaian yang berlaku untuk team atau jabatan"),
    ]
    add_table(doc, ("Istilah", "Definisi"), glossary_rows, [1.65, 5.25])

    configure_sections(doc)
    add_header_footer(doc)
    doc.core_properties.title = "Product Requirements Document Sistem Manajemen KPI PT Berlian Artha Label"
    doc.core_properties.subject = "Kebutuhan produk sistem KPI perusahaan, team, dan employee"
    doc.core_properties.keywords = "PRD, KPI, PT Berlian Artha Label, performance management"
    doc.core_properties.comments = "Draft untuk validasi bisnis"
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
